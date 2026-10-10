"""
خروجی Word — پورت مستقیم _generate_document_content از نسخهٔ اصلی Qt.
"""
import glob
import os
import random

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_TABLE_DIRECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import config


def _set_cell_shading(cell, fill_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_all_font_slots(rpr_owner_element, font_name):
    """نام فونت را برای همهٔ اسلات‌های Word ثبت می‌کند؛ متن فارسی از اسلات cs (Complex Script) می‌خواند."""
    rpr = rpr_owner_element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font_name)


def _apply_font_everywhere(doc, font_name):
    _set_all_font_slots(doc.styles["Normal"].element, font_name)
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            _set_all_font_slots(run._r, font_name)


def _pick_random_quote(filename):
    path = os.path.join(config.AKS_DIR, filename)
    if not os.path.exists(path):
        return ""
    try:
        df = pd.read_excel(path, header=None)
        data = df[0].dropna().tolist()
        return str(random.choice(data)) if data else ""
    except Exception:
        return ""


def generate_docx(schedule_state, motivational_sentences, font_name="Vazir"):
    student_name = schedule_state.student_name
    subjects = schedule_state.subjects
    days = schedule_state.days

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(12)
    style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for s in doc.styles:
        if hasattr(s, "paragraph_format"):
            p_pr = s.element.get_or_add_pPr()
            if p_pr.find(qn("w:bidi")) is None:
                p_pr.append(OxmlElement("w:bidi"))

    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Cm(38)
    section.page_height = Cm(21)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    image_paths = []
    for ext in ["*.png", "*.jpg", "*.jpeg", "*.gif", "*.bmp"]:
        image_paths.extend(glob.glob(os.path.join(config.AKS_DIR, ext)))
    if image_paths:
        try:
            doc.add_picture(random.choice(image_paths), width=Cm(15))
            doc.paragraphs[-1].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        except Exception:
            pass

    quote_s = _pick_random_quote("all-s.xlsx")
    if quote_s:
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.add_run(quote_s)
        run.font.name = font_name
        run.font.size = Pt(18)

    quote_l = _pick_random_quote("all-l.xlsx")
    if quote_l:
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.add_run(quote_l)
        run.font.name = font_name
        run.font.size = Pt(14)

    title = doc.add_heading(level=1)
    run = title.add_run(f"برنامه درسی {student_name}")
    run.font.size = Pt(18)
    run.font.name = font_name
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 77, 128)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    if motivational_sentences:
        sentence = random.choice(motivational_sentences).replace("{{name}}", student_name)
        p = doc.add_paragraph()
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.add_run(sentence)
        run.font.name = font_name
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 102, 204)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = subtitle.add_run("برنامه مطالعاتی هفتگی")
    run.font.name = font_name
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(77, 77, 77)
    doc.add_paragraph()

    num_cols = days + 1
    headers = [f"روز {i}" for i in range(days, 0, -1)] + ["درس"]
    table = doc.add_table(rows=len(subjects) + 1, cols=num_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.direction = WD_TABLE_DIRECTION.RTL

    header_color = "a2d2ff"
    row_colors = ["fefae0", "e9edc9"]

    for col, header_text in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header_text
        p = cell.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.runs[0]
        run.font.bold = True
        run.font.name = font_name
        run.font.size = Pt(13)
        _set_cell_shading(cell, header_color)

    for r, subject_data in enumerate(subjects, start=1):
        color = row_colors[(r - 1) % len(row_colors)]
        subject_cell = table.cell(r, num_cols - 1)
        p = subject_cell.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.add_run(f"{subject_data['subject']}\nساعت: {subject_data['hours']}")
        run.font.bold = True
        run.font.name = font_name
        _set_cell_shading(subject_cell, color)

        all_data = schedule_state.prepare_day_data(subject_data.get("data", []), days)
        manual = subject_data.get("manual")

        for col in range(days):
            cell_text = str(all_data[col])
            if manual and isinstance(manual, list) and col < len(manual) and manual[col]:
                cell_text = manual[col]
            else:
                extra = f"\nتعداد تست:{subject_data.get('tests','0')}\n \n \nتعداد تشریحی:{subject_data.get('descriptive','0')}\n"
                cell_text += extra

            day_cell = table.cell(r, days - 1 - col)
            day_cell.text = cell_text
            _set_cell_shading(day_cell, color)
            for paragraph in day_cell.paragraphs:
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                for run in paragraph.runs:
                    run.font.name = font_name

    _apply_font_everywhere(doc, font_name)
    return doc
